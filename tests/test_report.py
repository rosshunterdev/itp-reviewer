from src import report, schema

FINDINGS = [
    schema.Finding("Item 2.1", "acceptance_criteria", "Vague", "Client query", "consider a tolerance"),
    schema.Finding("Item 3.4", "hold_witness", "No hold point", "Risk", "consider adding a hold point"),
]

def test_markdown_groups_by_category_label():
    md = report.to_markdown(FINDINGS)
    assert "Acceptance Criteria" in md
    assert "Hold & Witness Points" in md
    assert "Item 2.1" in md and "consider a tolerance" in md

def test_markdown_empty():
    md = report.to_markdown([])
    assert "no findings" in md.lower()

def test_docx_returns_nonempty_bytes():
    data = report.to_docx(FINDINGS)
    assert isinstance(data, (bytes, bytearray)) and len(data) > 0


# ── run metadata header ───────────────────────────────────────────────

import io

from src import config, gen_report, gen_schema, prompts, report


def _meta():
    return report.build_meta("riverside_itp.xlsx", prompts.system_prompt())


def test_build_meta_records_model_prompt_hash_timestamp_source():
    m = _meta()
    assert m["model"] == config.MODEL
    assert m["source"] == "riverside_itp.xlsx"
    assert len(m["prompt_hash"]) == 10
    assert m["generated_at"][:2] == "20"  # ISO timestamp
    # hash changes when the prompt changes
    assert report.build_meta("x", "a")["prompt_hash"] != report.build_meta("x", "b")["prompt_hash"]


def test_review_markdown_header_has_all_four_fields():
    md = report.to_markdown([], _meta())
    for needle in ("riverside_itp.xlsx", config.MODEL, _meta()["prompt_hash"], "Generated"):
        assert needle in md
    assert "No findings were identified." in md


def test_review_markdown_without_meta_is_unchanged_shape():
    assert report.to_markdown([]).startswith("# ITP Review")
    assert "Model" not in report.to_markdown([])


def test_review_docx_header_has_meta():
    import docx
    data = report.to_docx([], _meta())
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    assert "riverside_itp.xlsx" in text and config.MODEL in text


def test_generated_itp_reports_carry_meta():
    import docx
    import openpyxl
    item = gen_schema.ITPItem(*["x"] * 6, "H", *["x"] * 3)
    m = _meta()
    assert m["source"] in gen_report.to_markdown([item], [], m)
    assert m["model"] in gen_report.to_markdown([], [], m)
    doc = docx.Document(io.BytesIO(gen_report.to_docx([item], [], m)))
    assert m["prompt_hash"] in "\n".join(p.text for p in doc.paragraphs)
    wb = openpyxl.load_workbook(io.BytesIO(gen_report.to_xlsx([item], [], m)))
    assert "Run info" in wb.sheetnames
    values = [c.value for row in wb["Run info"].iter_rows() for c in row]
    assert m["source"] in values and m["model"] in values
