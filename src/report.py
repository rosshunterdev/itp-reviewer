import hashlib
import io
from datetime import datetime
from src import config, schema


def build_meta(source: str, system_prompt: str) -> dict:
    """Run provenance for a report header. Call at run time, not at download time."""
    return {
        "source": source,
        "model": config.MODEL,
        "prompt_hash": hashlib.sha256(system_prompt.encode()).hexdigest()[:10],
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }


META_LABELS = [
    ("source", "Source file"),
    ("generated_at", "Generated"),
    ("model", "Model"),
    ("prompt_hash", "Prompt version (hash)"),
]


def meta_pairs(meta: dict | None) -> list[tuple[str, str]]:
    if not meta:
        return []
    return [(label, str(meta[key])) for key, label in META_LABELS if meta.get(key)]


def _group(findings):
    grouped = {c: [] for c in schema.CATEGORIES}
    for f in findings:
        grouped.get(f.category, grouped["other"]).append(f)
    return grouped


def to_markdown(findings, meta: dict | None = None) -> str:
    header = ["# ITP Review", ""]
    header += [f"- **{label}:** {value}" for label, value in meta_pairs(meta)]
    if meta:
        header.append("")
    if not findings:
        return "\n".join(header) + "\nNo findings were identified."
    lines = header + [f"{len(findings)} finding(s).", ""]
    grouped = _group(findings)
    for cat in schema.CATEGORIES:
        items = grouped[cat]
        if not items:
            continue
        lines.append(f"## {schema.CATEGORY_LABELS[cat]}")
        lines.append("")
        for f in items:
            lines.append(f"### {f.location}")
            lines.append(f"- **Observation:** {f.observation}")
            lines.append(f"- **Why it matters:** {f.why_it_matters}")
            lines.append(f"- **Suggested action:** {f.suggested_action}")
            lines.append("")
    return "\n".join(lines).strip() + "\n"


def to_docx(findings, meta: dict | None = None) -> bytes:
    import docx
    document = docx.Document()
    document.add_heading("ITP Review", level=0)
    for label, value in meta_pairs(meta):
        document.add_paragraph(f"{label}: {value}")
    if not findings:
        document.add_paragraph("No findings were identified.")
    else:
        document.add_paragraph(f"{len(findings)} finding(s).")
        grouped = _group(findings)
        for cat in schema.CATEGORIES:
            items = grouped[cat]
            if not items:
                continue
            document.add_heading(schema.CATEGORY_LABELS[cat], level=1)
            for f in items:
                document.add_heading(f.location, level=2)
                document.add_paragraph(f"Observation: {f.observation}")
                document.add_paragraph(f"Why it matters: {f.why_it_matters}")
                document.add_paragraph(f"Suggested action: {f.suggested_action}")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
