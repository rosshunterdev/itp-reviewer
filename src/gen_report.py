import io
from collections import OrderedDict
from src import gen_schema


COLUMNS = [
    ("Item", "item_number"),
    ("Work Package / Activity", "work_package"),
    ("Inspection / Test / Check", "inspection_test"),
    ("Acceptance Criteria", "acceptance_criteria"),
    ("Reference", "reference"),
    ("Frequency / Timing", "frequency"),
    ("Inspection Point", "inspection_point"),
    ("Contractor Responsibility", "contractor_resp"),
    ("Witness / Release By", "witness_release"),
    ("QA Record / Evidence", "qa_record"),
]


def _group_by_work_package(
    items: list[gen_schema.ITPItem],
) -> OrderedDict[str, list[gen_schema.ITPItem]]:
    groups: OrderedDict[str, list[gen_schema.ITPItem]] = OrderedDict()
    for item in items:
        groups.setdefault(item.work_package, []).append(item)
    return groups


def items_to_text(items: list[gen_schema.ITPItem]) -> str:
    if not items:
        return ""
    header = " | ".join(label for label, _ in COLUMNS)
    lines = [header]
    for item in items:
        row = " | ".join(getattr(item, field) for _, field in COLUMNS)
        lines.append(row)
    return "\n".join(lines)


def to_markdown(
    items: list[gen_schema.ITPItem],
    hold_points: list[gen_schema.HoldPoint],
) -> str:
    if not items:
        return "# Generated ITP\n\nNo ITP items were generated.\n"
    lines = [
        "# Generated ITP",
        "",
        f"{len(items)} item(s) across "
        f"{len(_group_by_work_package(items))} work package(s).",
        "",
    ]
    groups = _group_by_work_package(items)
    for wp, wp_items in groups.items():
        lines.append(f"## {wp}")
        lines.append("")
        for item in wp_items:
            lines.append(f"### {item.item_number} — {item.inspection_test}")
            lines.append(f"- **Acceptance Criteria:** {item.acceptance_criteria}")
            lines.append(f"- **Reference:** {item.reference}")
            lines.append(f"- **Frequency:** {item.frequency}")
            lines.append(f"- **Inspection Point:** {item.inspection_point}")
            lines.append(f"- **Contractor Responsibility:** {item.contractor_resp}")
            lines.append(f"- **Witness / Release By:** {item.witness_release}")
            lines.append(f"- **QA Record:** {item.qa_record}")
            lines.append("")
    if hold_points:
        lines.append("## Hold Point Register")
        lines.append("")
        for hp in hold_points:
            lines.append(
                f"- **{hp.hp_number}** (Item {hp.itp_item}): {hp.description}"
            )
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def to_docx(
    items: list[gen_schema.ITPItem],
    hold_points: list[gen_schema.HoldPoint],
) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Generated ITP", level=0)
    if not items:
        document.add_paragraph("No ITP items were generated.")
    else:
        document.add_paragraph(
            f"{len(items)} item(s) across "
            f"{len(_group_by_work_package(items))} work package(s)."
        )
        groups = _group_by_work_package(items)
        for wp, wp_items in groups.items():
            document.add_heading(wp, level=1)
            for item in wp_items:
                document.add_heading(
                    f"{item.item_number} — {item.inspection_test}", level=2
                )
                document.add_paragraph(
                    f"Acceptance Criteria: {item.acceptance_criteria}"
                )
                document.add_paragraph(f"Reference: {item.reference}")
                document.add_paragraph(f"Frequency: {item.frequency}")
                document.add_paragraph(f"Inspection Point: {item.inspection_point}")
                document.add_paragraph(
                    f"Contractor Responsibility: {item.contractor_resp}"
                )
                document.add_paragraph(
                    f"Witness / Release By: {item.witness_release}"
                )
                document.add_paragraph(f"QA Record: {item.qa_record}")
        if hold_points:
            document.add_heading("Hold Point Register", level=1)
            for hp in hold_points:
                document.add_paragraph(
                    f"{hp.hp_number} (Item {hp.itp_item}): {hp.description}"
                )
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def to_xlsx(
    items: list[gen_schema.ITPItem],
    hold_points: list[gen_schema.HoldPoint],
) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ITP"

    header_font = Font(bold=True, size=10)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, size=10, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = [label for label, _ in COLUMNS]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = wrap
        cell.border = thin_border

    for row_idx, item in enumerate(items, 2):
        for col_idx, (_, field) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=getattr(item, field))
            cell.alignment = wrap
            cell.border = thin_border

    col_widths = [8, 25, 30, 30, 25, 18, 14, 22, 20, 22]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    if hold_points:
        ws_hp = wb.create_sheet("Hold Point Register")
        hp_headers = ["HP Number", "ITP Item", "Description"]
        for col_idx, header in enumerate(hp_headers, 1):
            cell = ws_hp.cell(row=1, column=col_idx, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = wrap
            cell.border = thin_border
        for row_idx, hp in enumerate(hold_points, 2):
            ws_hp.cell(row=row_idx, column=1, value=hp.hp_number).border = thin_border
            ws_hp.cell(row=row_idx, column=2, value=hp.itp_item).border = thin_border
            cell = ws_hp.cell(row=row_idx, column=3, value=hp.description)
            cell.alignment = wrap
            cell.border = thin_border
        ws_hp.column_dimensions["A"].width = 12
        ws_hp.column_dimensions["B"].width = 12
        ws_hp.column_dimensions["C"].width = 50

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
