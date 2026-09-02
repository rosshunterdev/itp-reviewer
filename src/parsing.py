import io
import os
from dataclasses import dataclass, field


class UnsupportedFormat(Exception):
    pass


@dataclass
class ParsedDoc:
    text: str
    preview: str
    warnings: list[str] = field(default_factory=list)


def parse(filename: str, data: bytes) -> ParsedDoc:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".doc":
        raise UnsupportedFormat(
            "Legacy .doc files aren't supported. Please open it in Word and "
            "re-save as .docx, then upload again."
        )
    if ext == ".xlsx":
        return _parse_xlsx(data)
    if ext == ".pdf":
        return _parse_pdf(data)
    if ext == ".docx":
        return _parse_docx(data)
    raise UnsupportedFormat(
        f"Unsupported file type '{ext}'. Supported: .xlsx, .pdf, .docx."
    )


def _row_to_line(cells) -> str:
    vals = ["" if c is None else str(c).strip() for c in cells]
    if not any(vals):
        return ""
    return " | ".join(vals)


def _parse_xlsx(data: bytes) -> ParsedDoc:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"### Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            line = _row_to_line(row)
            if line:
                lines.append(line)
        lines.append("")
    text = "\n".join(lines).strip()
    return ParsedDoc(text=text, preview=text)


def _parse_pdf(data: bytes) -> ParsedDoc:
    import pdfplumber
    lines, warnings = [], []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            lines.append(f"### Page {i}")
            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    line = _row_to_line(row)
                    if line:
                        lines.append(line)
            page_text = page.extract_text() or ""
            if page_text.strip():
                lines.append(page_text.strip())
            lines.append("")
    text = "\n".join(lines).strip()
    if not text:
        warnings.append("No extractable text found — the PDF may be scanned images.")
    return ParsedDoc(text=text, preview=text, warnings=warnings)


def _parse_docx(data: bytes) -> ParsedDoc:
    import docx
    document = docx.Document(io.BytesIO(data))
    lines = []
    for para in document.paragraphs:
        if para.text.strip():
            lines.append(para.text.strip())
    for t_i, table in enumerate(document.tables, 1):
        lines.append(f"### Table {t_i}")
        for row in table.rows:
            line = _row_to_line([c.text for c in row.cells])
            if line:
                lines.append(line)
    text = "\n".join(lines).strip()
    return ParsedDoc(text=text, preview=text)
